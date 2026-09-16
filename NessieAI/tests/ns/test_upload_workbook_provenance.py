import json
import openpyxl

from nextseek_api.batch_upload.convert import detect_format, parse_traditional_file
from NessieAI.ns.upload_workbook import MODE_UPDATE, render_upload_workbook

ROWS = [{"json_metadata": {"UID": "D.SEQ-EXAMPLE-1", "MappedPercent": 91.4,
                           "ContamPercent": 3.2}, "assay_ids": []}]
PROV = [
    {"uid": "D.SEQ-EXAMPLE-1", "attribute": "MappedPercent", "value": 91.4,
     "origin": "map", "raw_key": "STAR_uniquely_mapped_percent",
     "source_file": "multiqc/star_salmon/multiqc_data/multiqc_general_stats.txt"},
    {"uid": "D.SEQ-EXAMPLE-1", "attribute": "ContamPercent", "value": 3.2,
     "origin": "proposed", "raw_key": "Kraken2_bracken_fraction_total_reads",
     "source_file": "star_salmon/contaminants/kraken2/x.kraken2.report.txt"},
]


def _render(tmp_path):
    out = tmp_path / "prov.xlsx"
    render_upload_workbook("D.SEQ", ROWS, str(out), mode=MODE_UPDATE, provenance=PROV)
    return out


def test_provenance_sheet_has_one_row_per_entry(tmp_path):
    sheet = openpyxl.load_workbook(_render(tmp_path))["Provenance"]
    assert sheet.max_row == len(PROV) + 1
    assert [c.value for c in sheet[1]] == [
        "UID", "Attribute", "Value", "Origin", "Raw key", "Source file"]


def test_an_unapproved_cell_is_visibly_marked_in_the_sheet(tmp_path):
    sheet = openpyxl.load_workbook(_render(tmp_path))["Provenance"]
    origins = {sheet.cell(row=r, column=2).value: sheet.cell(row=r, column=4).value
               for r in range(2, sheet.max_row + 1)}
    assert origins["ContamPercent"] == "proposed"
    assert origins["MappedPercent"] == "map"


def test_the_fifth_sheet_does_not_break_format_detection(tmp_path):
    assert detect_format(str(_render(tmp_path))) == "traditional"


def test_the_fifth_sheet_produces_no_dropped_column_warning(tmp_path):
    batch = parse_traditional_file(str(_render(tmp_path)))
    joined = " ".join(batch.warnings or [])
    assert "Provenance" not in joined
    assert json.loads(batch.rows[0].json_metadata)["MappedPercent"] == 91.4


def test_no_provenance_means_no_fifth_sheet(tmp_path):
    out = tmp_path / "bare.xlsx"
    render_upload_workbook("D.SEQ", ROWS, str(out), mode=MODE_UPDATE)
    assert "Provenance" not in openpyxl.load_workbook(out).sheetnames
