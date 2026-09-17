import json

import openpyxl
import pytest

from nextseek_api.batch_upload.convert import parse_traditional_file
from NessieAI.ns.upload_workbook import MODE_NEW, MODE_UPDATE, render_upload_workbook

NEW_ROWS = [{"json_metadata": {"Parent": "D.SEQ-EXAMPLE-1", "Scientist": "A Person"},
             "assay_ids": [7]}]
UPD_ROWS = [{"json_metadata": {"UID": "D.SEQ-EXAMPLE-1", "MappedPercent": 91.4},
             "assay_ids": []},
            {"json_metadata": {"UID": "D.SEQ-EXAMPLE-2", "MappedPercent": 90.7},
             "assay_ids": []}]
BOOL_ROWS = [{"json_metadata": {"Parent": "D.SEQ-EXAMPLE-1", "Flag": True}, "assay_ids": []},
             {"json_metadata": {"Parent": "D.SEQ-EXAMPLE-2", "Flag": False}, "assay_ids": []}]
MIXED_ROWS = [{"json_metadata": {"Parent": "D.SEQ-EXAMPLE-1", "Mixed": 5}, "assay_ids": []},
              {"json_metadata": {"Parent": "D.SEQ-EXAMPLE-2", "Mixed": "n/a"}, "assay_ids": []}]
BLANK_ROWS = [{"json_metadata": {"Parent": "D.SEQ-EXAMPLE-1", "Blank": None}, "assay_ids": []},
              {"json_metadata": {"Parent": "D.SEQ-EXAMPLE-2", "Blank": ""}, "assay_ids": []}]
PROVENANCE_ROWS = [{"uid": "D.SEQ-EXAMPLE-1", "attribute": "MappedPercent", "value": 91.4,
                     "origin": "harvest", "raw_key": "mapped_pct", "source_file": "run.json"}]


def _field_types(out):
    """{Field: Field Type} as declared in the Instructions sheet of a rendered workbook."""
    sheet = openpyxl.load_workbook(out)["Instructions"]
    return {row[0].value: row[2].value for row in sheet.iter_rows(min_row=2)}


def test_new_mode_emits_no_uid_column(tmp_path):
    out = tmp_path / "new.xlsx"
    render_upload_workbook("A.GEX", NEW_ROWS, str(out), mode=MODE_NEW)
    headers = [c.value for c in openpyxl.load_workbook(out)["Samples"][1]]
    assert "UID" not in headers


def test_update_mode_puts_uid_first_and_populated(tmp_path):
    out = tmp_path / "upd.xlsx"
    render_upload_workbook("D.SEQ", UPD_ROWS, str(out), mode=MODE_UPDATE)
    sheet = openpyxl.load_workbook(out)["Samples"]
    assert [c.value for c in sheet[1]][0] == "UID"
    assert sheet.cell(row=2, column=1).value == "D.SEQ-EXAMPLE-1"


def test_update_mode_does_not_declare_uid_in_instructions(tmp_path):
    out = tmp_path / "upd.xlsx"
    render_upload_workbook("D.SEQ", UPD_ROWS, str(out), mode=MODE_UPDATE)
    fields = [r[0].value for r in openpyxl.load_workbook(out)["Instructions"].iter_rows(min_row=2)]
    assert "UID" not in fields


def test_update_mode_round_trips_with_the_uid_preserved(tmp_path):
    out = tmp_path / "upd.xlsx"
    render_upload_workbook("D.SEQ", UPD_ROWS, str(out), mode=MODE_UPDATE)
    batch = parse_traditional_file(str(out))
    assert [r.UID for r in batch.rows] == ["D.SEQ-EXAMPLE-1", "D.SEQ-EXAMPLE-2"]
    # json_metadata is InputRowModel's minified-JSON *string* field (see
    # nextseek_api/batch_upload/models.py:119 and the pinned round-trip test's
    # own json.loads(r0.json_metadata) usage) — not a dict, so it must be
    # parsed before indexing.
    assert json.loads(batch.rows[0].json_metadata)["MappedPercent"] == 91.4


def test_update_mode_rejects_a_row_with_no_uid(tmp_path):
    with pytest.raises(ValueError, match="UID"):
        render_upload_workbook("D.SEQ", [{"json_metadata": {"MappedPercent": 1}}],
                               str(tmp_path / "x.xlsx"), mode=MODE_UPDATE)


def test_new_mode_rejects_a_row_carrying_a_uid(tmp_path):
    # Symmetric to test_update_mode_rejects_a_row_with_no_uid: new mode has
    # no UID column at all, so a row carrying one must raise loudly instead
    # of having its UID silently stripped by the field-set loop -- which
    # would otherwise round-trip through parse_traditional_file as a
    # brand-new sample, duplicating whatever it was meant to update.
    with pytest.raises(ValueError, match="UID"):
        render_upload_workbook(
            "D.SEQ",
            [{"json_metadata": {"UID": "D.SEQ-EXAMPLE-1", "Parent": "D.SEQ-EXAMPLE-1"}}],
            str(tmp_path / "x.xlsx"), mode=MODE_NEW)


def test_all_four_sheets_are_present_in_both_modes(tmp_path):
    for mode, rows, stype in ((MODE_NEW, NEW_ROWS, "A.GEX"), (MODE_UPDATE, UPD_ROWS, "D.SEQ")):
        out = tmp_path / f"{mode}.xlsx"
        render_upload_workbook(stype, rows, str(out), mode=mode)
        names = openpyxl.load_workbook(out).sheetnames
        assert {"Instructions", "Samples", "Assay", "Ontology"} <= set(names)


def test_unknown_mode_raises_value_error(tmp_path):
    with pytest.raises(ValueError, match="unknown mode"):
        render_upload_workbook("A.GEX", NEW_ROWS, str(tmp_path / "x.xlsx"), mode="bogus")


def test_provenance_sheet_has_the_expected_header_row(tmp_path):
    out = tmp_path / "prov.xlsx"
    render_upload_workbook("D.SEQ", UPD_ROWS, str(out), mode=MODE_UPDATE,
                           provenance=PROVENANCE_ROWS)
    header = [c.value for c in openpyxl.load_workbook(out)["Provenance"][1]]
    assert header == ["UID", "Attribute", "Value", "Origin", "Raw key", "Source file"]


def test_field_type_declares_number_for_an_all_numeric_field(tmp_path):
    out = tmp_path / "upd.xlsx"
    render_upload_workbook("D.SEQ", UPD_ROWS, str(out), mode=MODE_UPDATE)
    assert _field_types(out)["MappedPercent"] == "Number"


def test_field_type_declares_text_for_a_text_field(tmp_path):
    out = tmp_path / "new.xlsx"
    render_upload_workbook("A.GEX", NEW_ROWS, str(out), mode=MODE_NEW)
    assert _field_types(out)["Scientist"] == "Text"


def test_field_type_bool_carve_out_declares_text_not_number(tmp_path):
    # Python bools are ints; a naive isinstance(value, (int, float)) check
    # would misclassify True/False as "Number".
    out = tmp_path / "bool.xlsx"
    render_upload_workbook("A.GEX", BOOL_ROWS, str(out), mode=MODE_NEW)
    assert _field_types(out)["Flag"] == "Text"


def test_field_type_mixed_numeric_and_text_declares_text(tmp_path):
    out = tmp_path / "mixed.xlsx"
    render_upload_workbook("A.GEX", MIXED_ROWS, str(out), mode=MODE_NEW)
    assert _field_types(out)["Mixed"] == "Text"


def test_field_type_all_blank_field_falls_back_to_text(tmp_path):
    out = tmp_path / "blank.xlsx"
    render_upload_workbook("A.GEX", BLANK_ROWS, str(out), mode=MODE_NEW)
    assert _field_types(out)["Blank"] == "Text"
