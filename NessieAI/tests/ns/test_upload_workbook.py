"""render_upload_workbook must produce a 4-sheet workbook that round-trips cleanly
through the real batch-upload parser (the correctness oracle)."""
import json

from NessieAI.ns.upload_workbook import render_upload_workbook
from nextseek_api.batch_upload.convert import detect_format, parse_traditional_file


def test_render_round_trips_through_parse_traditional(tmp_path):
    rows = [
        {"json_metadata": {"Parent": "D.SEQ-220823SHA-1", "Scientist": "Marie Floryan",
                           "File_PrimaryData": "gideon_1_matrix.h5"}, "assay_ids": [12]},
        {"json_metadata": {"Parent": "D.SEQ-220823SHA-2", "Scientist": "Marie Floryan",
                           "File_PrimaryData": "gideon_2_matrix.h5"}, "assay_ids": [12]},
    ]
    out = str(tmp_path / "reingest_A.SCXP.xlsx")
    render_upload_workbook("A.SCXP", rows, out)

    # It must be recognized as the 4-sheet (traditional) format, not flat.
    assert detect_format(out) == "traditional"

    batch = parse_traditional_file(out)
    assert len(batch.rows) == 2

    r0 = batch.rows[0]
    assert r0.SampleType == "A.SCXP"
    assert not r0.UID  # blank -> server mints
    meta = json.loads(r0.json_metadata)
    assert meta["Parent"] == "D.SEQ-220823SHA-1"
    assert meta["Scientist"] == "Marie Floryan"
    assert meta["File_PrimaryData"] == "gideon_1_matrix.h5"
    assert r0.assay_ids == [12]

    r1 = batch.rows[1]
    assert json.loads(r1.json_metadata)["Parent"] == "D.SEQ-220823SHA-2"


def test_render_multi_parent_and_no_assay(tmp_path):
    # A merged/aggregate output derived from several inputs (;-delimited Parent),
    # and no assay_ids (ASSAY sheet header-only -> assay_ids == []).
    rows = [
        {"json_metadata": {"Parent": "D.SEQ-220823SHA-1;D.SEQ-220823SHA-2",
                           "Scientist": "Marie Floryan"}, "assay_ids": []},
    ]
    out = str(tmp_path / "reingest_A.SCXP_merged.xlsx")
    render_upload_workbook("A.SCXP", rows, out)

    batch = parse_traditional_file(out)
    assert len(batch.rows) == 1
    meta = json.loads(batch.rows[0].json_metadata)
    assert meta["Parent"] == "D.SEQ-220823SHA-1;D.SEQ-220823SHA-2"
    assert batch.rows[0].assay_ids == []


def test_render_rejects_empty_rows(tmp_path):
    import pytest
    with pytest.raises(ValueError):
        render_upload_workbook("A.SCXP", [], str(tmp_path / "empty.xlsx"))


def test_render_keeps_blank_schema_attributes_as_columns(tmp_path):
    # Every schema attribute gets a column even when no row has a value for
    # it, in schema order, with extra row keys after; the blanks never reach
    # json_metadata, so an update upload cannot clear an existing value.
    import openpyxl

    rows = [{"json_metadata": {"Parent": "D.SEQ-220823SHA-1", "Extra": "x"},
             "assay_ids": []}]
    out = str(tmp_path / "reingest_A.SCXP_blanks.xlsx")
    render_upload_workbook("A.SCXP", rows, out,
                           schema_fields=["UID", "Name", "Parent", "Software"])

    wb = openpyxl.load_workbook(out)
    header = [c.value for c in wb["Samples"][1]]
    assert header == ["Name", "Parent", "Software", "Extra"]
    declared = [r[0] for r in wb["Instructions"].iter_rows(min_row=2, values_only=True)]
    assert declared == header

    meta = json.loads(parse_traditional_file(out).rows[0].json_metadata)
    assert meta == {"Parent": "D.SEQ-220823SHA-1", "Extra": "x"}


def test_render_rejects_rows_with_no_metadata_despite_schema(tmp_path):
    import pytest
    with pytest.raises(ValueError, match="no json_metadata"):
        render_upload_workbook("A.SCXP", [{"json_metadata": {}}],
                               str(tmp_path / "blank.xlsx"), schema_fields=["Name"])


def test_formula_shaped_value_is_a_text_cell(tmp_path):
    from openpyxl import load_workbook

    evil = '=HYPERLINK("http://x","y")'
    out = tmp_path / "u.xlsx"
    render_upload_workbook("MUS", [{"json_metadata": {"Name": evil}, "assay_ids": []}], str(out))
    cell = load_workbook(out)["Samples"]["A2"]
    assert cell.data_type == "s" and cell.value == evil


def test_every_user_value_is_a_text_cell(tmp_path):
    # SEC-H03 across the whole workbook: a field name (Instructions, the
    # Samples header), an update-mode UID and a Provenance value all come from
    # stored metadata or a curator's chat answer, so none may become a formula.
    from openpyxl import load_workbook

    evil = '=HYPERLINK("http://x","y")'
    out = tmp_path / "u.xlsx"
    render_upload_workbook(
        "MUS", [{"json_metadata": {"UID": evil, evil: "v"}, "assay_ids": []}], str(out),
        mode="update", provenance=[{"uid": evil, "attribute": evil, "value": evil,
                                    "origin": evil, "raw_key": evil,
                                    "source_file": evil, "answered_by": evil}])
    wb = load_workbook(out)
    formulas = [(ws.title, c.coordinate) for ws in wb.worksheets
                for row in ws.iter_rows() for c in row if c.data_type == "f"]
    assert formulas == []
