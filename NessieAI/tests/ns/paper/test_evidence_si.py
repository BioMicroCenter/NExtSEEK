import io
import warnings
import zipfile

from openpyxl import Workbook

from NessieAI.ns.paper.evidence import deposit_items, paper_items, record_evidence_ids
from NessieAI.ns.paper.jats import parse_jats
from NessieAI.ns.paper.models import DepositRecord, FileRef, Gap, SupplementRef
from NessieAI.ns.paper.si import supplement_items
from NessieAI.tests.ns.paper.fixtures import JATS_XML


def _zip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


def _xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Samples"
    ws.append(["Mouse", "Sex"])
    ws.append(["M1", "female"])
    ws["D9"] = None
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_paper_items_cover_paragraphs_and_captions():
    ids = [i.id for i in paper_items(parse_jats(JATS_XML))]
    assert "E:methods:p1" in ids and "E:cap:1" in ids


def test_deposit_items_and_ids():
    rec = DepositRecord(repo="geo", record_id="GSM1", kind="sample", title="t1",
                        fields={"characteristics.sex": "female"},
                        files=[FileRef(name="a.gz", url="ftp://x/a.gz", size=3, checksum="ab", checksum_type="md5")])
    items = {i.id: i.text for i in deposit_items(rec)}
    assert items == {"E:geo:GSM1:characteristics.sex": "female", "E:geo:GSM1:title": "t1",
                     "E:file:geo:GSM1:a.gz": "a.gz ftp://x/a.gz size=3 md5=ab"}
    assert record_evidence_ids(rec) == list(items)


def test_duplicate_file_names_in_one_record_get_distinct_ids():
    rec = DepositRecord(repo="geo", record_id="GSM1", kind="sample", files=[
        FileRef(name="a.gz", url="ftp://x/1/a.gz"), FileRef(name="a.gz", url="ftp://x/2/a.gz")])
    items = {i.id: i.text for i in deposit_items(rec)}
    assert items == {"E:file:geo:GSM1:a.gz": "a.gz ftp://x/1/a.gz",
                     "E:file:geo:GSM1:a.gz#2": "a.gz ftp://x/2/a.gz"}
    assert record_evidence_ids(rec) == list(items)


def test_supplement_items_reads_tables_and_reports_gaps():
    data = _zip({"suppl/NIHMS1_TableS1.xlsx": _xlsx(), "TableS2.csv": b"id,dose\nM1,5 mg\n",
                 "methods.docx": b"PK..", "notes.txt": b"First note.\n\nSecond note.\n"})
    items, gaps = supplement_items(data, [SupplementRef(name="NIHMS1_TableS1.xlsx"), SupplementRef(name="MovieS1.mp4")])
    got = {i.id: i.text for i in items}
    assert got["E:si:NIHMS1_TableS1.xlsx!Samples!B2"] == "female"
    assert got["E:si:TableS2.csv!B2"] == "5 mg"
    assert got["E:si:notes.txt!p2"] == "Second note."
    assert {(g.kind, g.detail.split(":")[0]) for g in gaps} == {
        ("supplement_unparsed", "methods.docx"), ("supplement_unavailable", "MovieS1.mp4")}


def test_same_basename_in_two_folders_keeps_ids_unique():
    data = _zip({"a/TableS1.csv": b"first\n", "b/TableS1.csv": b"second\n"})
    items, _ = supplement_items(data, [])
    got = {i.id: i.text for i in items}
    assert got == {"E:si:TableS1.csv!A1": "first", "E:si:b/TableS1.csv!A1": "second"}


def test_unreadable_archive_becomes_gaps_not_a_failure():
    items, gaps = supplement_items(b"not a zip", [SupplementRef(name="TableS1.xlsx")])
    assert items == []
    assert gaps == [Gap(kind="supplement_unparsed", detail="supplementary archive: BadZipFile"),
                    Gap(kind="supplement_unavailable", detail="TableS1.xlsx")]


def test_corrupt_member_is_a_gap_naming_the_file():
    items, gaps = supplement_items(_zip({"bad.xlsx": b"garbage bytes"}), [])
    assert items == []
    assert [(g.kind, g.detail.split(":")[0]) for g in gaps] == [("supplement_unparsed", "bad.xlsx")]


def test_empty_readable_file_is_a_gap():
    items, gaps = supplement_items(_zip({"empty.csv": b""}), [])
    assert items == []
    assert gaps == [Gap(kind="supplement_unparsed", detail="empty.csv: no readable text")]


def test_tsv_cells_get_cell_ids():
    items, gaps = supplement_items(_zip({"TableS3.tsv": b"id\tdose\nM2\t7 mg\n"}), [])
    assert gaps == []
    assert {i.id: i.text for i in items}["E:si:TableS3.tsv!B2"] == "7 mg"


def test_member_order_never_duplicates_a_stem():
    items, _ = supplement_items(_zip({"a/T.csv": b"first\n", "T.csv": b"second\n"}), [])
    assert sorted((i.id, i.text) for i in items) == [("E:si:T.csv!A1", "first"), ("E:si:T.csv#2!A1", "second")]


def test_exact_duplicate_path_stays_unique():
    buf = io.BytesIO()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)  # zipfile warns on a repeated name; the archive is the test input
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("a/T.csv", b"one\n")
            zf.writestr("a/T.csv", b"two\n")
    items, _ = supplement_items(buf.getvalue(), [])
    assert sorted(i.id for i in items) == ["E:si:T.csv!A1", "E:si:a/T.csv!A1"]


# Cells read as a curator sees them, so a sample name matches verbatim.

def _formatted_xlsx() -> bytes:
    from datetime import datetime, time

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    cells = {"A1": (1, None), "A2": (1.0, None), "A3": (0.1 + 0.2, None), "A4": (1e-07, None),
             "A5": (datetime(2024, 2, 3), "yyyy-mm-dd"), "A6": (datetime(2024, 2, 3, 10, 30), "yyyy-mm-dd h:mm"),
             "A7": (time(10, 30), "h:mm"), "A8": (True, None), "A9": (0.25, "0%"), "A10": (0.125, "0.0%"),
             "A11": (7, "000"), "A12": (1.5, "0.00"), "A13": (2.675, "0.00"), "A14": ("  M1 ", None),
             "A15": (123456789012, None), "A16": (-3.0, None)}
    for coord, (value, fmt) in cells.items():
        ws[coord] = value
        if fmt:
            ws[coord].number_format = fmt
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_workbook_cells_read_as_a_curator_sees_them():
    items, gaps = supplement_items(_zip({"T.xlsx": _formatted_xlsx()}), [])
    assert gaps == []
    got = {i.id.rsplit("!", 1)[1]: i.text for i in items}
    assert got == {"A1": "1", "A2": "1", "A3": "0.3", "A4": "1E-07", "A5": "2024-02-03", "A6": "2024-02-03 10:30",
                   "A7": "10:30", "A8": "TRUE", "A9": "25%", "A10": "12.5%", "A11": "007", "A12": "1.50",
                   "A13": "2.68", "A14": "M1", "A15": "123456789012", "A16": "-3"}


# Decompression is bounded: a workbook is itself a zip, and the archive as a whole has a budget.

def _padded_xlsx(pad_bytes: int) -> bytes:
    """A real workbook plus one highly compressible member, as a bomb would carry."""
    src = zipfile.ZipFile(io.BytesIO(_xlsx()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            out.writestr(info, src.read(info))
        out.writestr("xl/padding.xml", b"\0" * pad_bytes)
    return buf.getvalue()


def test_workbook_that_unpacks_past_the_cap_is_a_gap_and_never_opened(monkeypatch):
    import NessieAI.ns.paper.si as si

    monkeypatch.setattr(si, "MAX_WORKBOOK_UNPACKED_BYTES", 200_000)
    opened = []
    monkeypatch.setattr(si, "_workbook", lambda name, data: opened.append(name) or [])
    bomb = _padded_xlsx(1_000_000)
    assert len(bomb) < 200_000  # small on the wire, large unpacked
    items, gaps = supplement_items(_zip({"Bomb.xlsx": bomb}), [])
    assert items == [] and opened == []
    assert [(g.kind, g.detail.split(":")[0]) for g in gaps] == [("supplement_unparsed", "Bomb.xlsx")]
    assert "unpacked" in gaps[0].detail


def test_workbook_under_the_cap_is_still_read(monkeypatch):
    import NessieAI.ns.paper.si as si

    monkeypatch.setattr(si, "MAX_WORKBOOK_UNPACKED_BYTES", 2_000_000)
    items, gaps = supplement_items(_zip({"T.xlsx": _padded_xlsx(1_000_000)}), [])
    assert gaps == [] and {i.text for i in items} >= {"M1", "female"}


def test_archive_unpacked_budget_turns_the_rest_into_gaps(monkeypatch):
    import NessieAI.ns.paper.si as si

    monkeypatch.setattr(si, "MAX_ARCHIVE_UNPACKED_BYTES", 2_500)
    rows = b"id\n" + b"M1\n" * 500  # 1503 bytes unpacked each
    items, gaps = supplement_items(_zip({"A.csv": rows, "B.csv": rows, "C.txt": b"short note"}), [])
    assert {i.id.split("!")[0] for i in items} == {"E:si:A.csv", "E:si:C.txt"}  # a member that still fits is read
    assert [(g.kind, g.detail.split(":")[0]) for g in gaps] == [("supplement_unparsed", "B.csv")]
    assert "budget" in gaps[0].detail


# Europe PMC's zip also carries the article's own figure images (gr1.gif, gr1.jpg, nihms-...-f0004.gif):
# they are read as captions, so they are not supplements, and not gaps.

def test_article_figure_images_in_the_zip_are_not_gaps():
    data = _zip({"gr1.gif": b"GIF89a", "gr1.jpg": b"\xff\xd8", "nihms-1823260-f0004.gif": b"GIF89a",
                 "mmc1.docx": b"PK..", "TableS1.csv": b"a\n"})
    items, gaps = supplement_items(data, [SupplementRef(name="mmc1.docx"), SupplementRef(name="TableS1.csv")])
    assert [i.id for i in items] == ["E:si:TableS1.csv!A1"]
    assert gaps == [Gap(kind="supplement_unparsed", detail="mmc1.docx: .docx is not read")]


def test_an_image_the_paper_lists_as_a_supplement_is_still_a_gap():
    items, gaps = supplement_items(_zip({"suppl/FigS1.TIF": b"II*\0"}), [SupplementRef(name="FigS1.tif")])
    assert items == [] and gaps == [Gap(kind="supplement_unparsed", detail="FigS1.TIF: .tif is not read")]


def test_listed_names_match_members_across_folder_case_and_url_encoding():
    data = _zip({"suppl/Table S1.CSV": b"a\n", "SUPP/data-s2.TSV": b"b\n"})
    items, gaps = supplement_items(data, [SupplementRef(name="Table S1.csv"), SupplementRef(name="Data-S2.tsv"),
                                          SupplementRef(name="Table%20S1.csv")])
    assert gaps == [] and len(items) == 2


def test_a_supplement_listed_twice_is_one_gap():
    _, gaps = supplement_items(_zip({"x.csv": b"a\n"}), [SupplementRef(name="mmc9.pdf"), SupplementRef(name="mmc9.pdf")])
    assert gaps == [Gap(kind="supplement_unavailable", detail="mmc9.pdf")]
