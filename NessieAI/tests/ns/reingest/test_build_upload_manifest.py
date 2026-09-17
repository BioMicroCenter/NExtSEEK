"""granular._build_upload_xlsx (Task 7): manifest_id, not values.

CC sends only ``manifest_id`` + ``mode``; the server loads its own copy of the
manifest (``store.load_manifest``), maps it (``mapper.apply``), and derives
every row itself -- a measured value never round-trips through the model.
Legacy ``rows``-based calls (the pre-Task-7 shape) keep working unchanged.

Openpyxl is used directly (not ``parse_traditional_file``) to read the
Provenance sheet: the traditional parser deliberately never reads it (see
``NessieAI/ns/CLAUDE.md``'s "An optional fifth Provenance sheet is inert to
the parser").
"""
from __future__ import annotations

import inspect
from unittest.mock import patch

import openpyxl
import pytest

import NessieAI.ns.granular as g
from NessieAI.ns.reingest import manifest as manifest_mod
from NessieAI.ns.reingest import store as store_mod

pytestmark = pytest.mark.django_db


def _meta(row) -> dict:
    """A parsed InputRowModel's json_metadata is a JSON STRING
    (nextseek_api/batch_upload/models.py), not a dict -- decode it."""
    import json
    return json.loads(row.json_metadata)


def _dispatch(op: str, args: dict, config=None, session=None, write_gate=None,
              neo4j_exec=None, outputs_dir=None):
    """The brief's sketch calls a `granular.dispatch(...)` that does not exist
    (the real entry point is `run_op`, keyword-only past `args`; see
    `NessieAI/tests/ns/reingest/test_run_harvest_op.py`'s identical note).
    This wrapper keeps the brief's positional call shape at the test call
    sites below while calling the real function correctly."""
    return g.run_op(op, args, config=config, session=session, write_gate=write_gate,
                     neo4j_exec=neo4j_exec, outputs_dir=outputs_dir)


# A D.SEQ catalog row that declares MappedPercent (and Genome, which the
# rnaseq map's qc_attributes ALSO backfills unconditionally from
# $params.genome, present or not) as real attributes, but NOT rRNAPercent --
# so a manifest carrying all three metrics exercises both branches of the
# attribute_exists parking decision in one run, with Genome kept boring
# (always resolvable, always declared) so it never contaminates a test's
# parked-attribute count. Shape mirrors
# NessieAI/tests/ns/reingest/test_proposals.py's SAMPLE_TYPE_ROW fixture (the
# committed pattern for exercising the real context_catalog loader without a
# seeded database).
_D_SEQ_ROW = {
    "sample_type": "D.SEQ", "sampletype_id": 1, "name": "Sequencing Data",
    "description": "Raw sequencing reads.", "clade": "Raw", "tags": "",
    "required_metadata": "UID",
    "standard_metadata": "MappedPercent, Genome",
    "possible_metadata_fields": "",
    "parent_sampletypes": "", "child_sampletypes": "",
    "associated_assay_parents": "", "associated_assay_children": "",
}


def _save_manifest(tmp_path, monkeypatch, *, metrics=None):
    # store._ROOT is resolved once at module-import time from an env var; a
    # module already imported by an earlier test would ignore a later
    # monkeypatch.setenv, so patch the module attribute directly instead
    # (same pattern as test_run_harvest_op.py's _patch_harvest).
    monkeypatch.setattr(store_mod, "_ROOT", str(tmp_path / "manifests"))
    run_manifest = manifest_mod.RunManifest(
        run_dir="/net/cluster/runs/r1",
        pipeline=manifest_mod.PipelineInfo(
            name="nf-core/rnaseq", version="3.18.0", run_name="test_run_1"),
        params={"genome": "GRCh38", "aligner": "star_salmon"},
        samples=[manifest_mod.SampleRecord(
            nfcore_sample="SAMPLE_1", d_seq_uid="D.SEQ-EXAMPLE-1",
            uid_resolution=manifest_mod.RESOLUTION_LAUNCH_RECORD,
            metrics=metrics or {})],
        sources={"metrics": "multiqc/star_salmon/multiqc_data/multiqc_general_stats.txt",
                 "params": "params.json"},
    )
    return store_mod.save_manifest(run_manifest)


# ---------------------------------------------------------------------------
# Step 1 (brief): validation
# ---------------------------------------------------------------------------

def test_a_bad_manifest_id_is_rejected_without_touching_disk():
    with pytest.raises(g.OpValidationError):
        _dispatch("build-upload-xlsx", {"manifest_id": "../../etc/passwd", "mode": "update"})


def test_an_unknown_mode_is_rejected():
    with pytest.raises(g.OpValidationError):
        _dispatch("build-upload-xlsx", {"manifest_id": "abc123", "mode": "sideways"})


def test_neither_manifest_id_nor_rows_is_rejected():
    with pytest.raises(g.OpValidationError):
        _dispatch("build-upload-xlsx", {})


def test_the_op_returns_proposals_rather_than_writing_them():
    """build-upload-xlsx never writes to NExtSEEK. Proposal rows are persisted by
    the service layer, which already persists the artifact bundle."""
    source = inspect.getsource(g._build_upload_xlsx)
    assert "proposals.record" not in source
    assert "proposals" in source  # returned in the envelope


def test_the_manifest_path_returns_proposals_rather_than_writing_them():
    """The dispatcher above is a thin manifest_id/rows router (see
    _build_upload_xlsx's own docstring); the real invariant belongs on the
    manifest-driven implementation, so it is checked there too."""
    source = inspect.getsource(g._build_upload_xlsx_from_manifest)
    assert "proposals.record" not in source
    assert "proposals" in source


# ---------------------------------------------------------------------------
# mode=new: analysis children, never the D.SEQ backfill
# ---------------------------------------------------------------------------

def test_mode_new_renders_analysis_children_not_dseq(tmp_path, monkeypatch):
    manifest_id = _save_manifest(tmp_path, monkeypatch)
    result = _dispatch("build-upload-xlsx", {"manifest_id": manifest_id, "mode": "new"},
                        outputs_dir=str(tmp_path))

    assert set(result["saved_files"]) == {"reingest_A_ALN", "reingest_A_GEX"}
    assert result["qa"]["A.ALN"]["disposition"] in ("CLEAN", "SOFT_FLAG")
    assert result["qa"]["A.GEX"]["disposition"] in ("CLEAN", "SOFT_FLAG")
    assert "D.SEQ" not in result["qa"]
    assert result["reply"]  # rendered by report.render_qa_for_user, non-empty
    assert result["proposals"] == []  # every metric in this manifest is mapped


def test_mode_new_a_dot_aln_workbook_round_trips(tmp_path, monkeypatch):
    manifest_id = _save_manifest(tmp_path, monkeypatch)
    result = _dispatch("build-upload-xlsx", {"manifest_id": manifest_id, "mode": "new"},
                        outputs_dir=str(tmp_path))

    from nextseek_api.batch_upload.convert import parse_traditional_file
    batch = parse_traditional_file(result["saved_files"]["reingest_A_ALN"])
    assert len(batch.rows) == 1
    assert batch.rows[0].SampleType == "A.ALN"
    assert _meta(batch.rows[0])["Parent"] == "D.SEQ-EXAMPLE-1"


# ---------------------------------------------------------------------------
# mode=update: the D.SEQ backfill, attribute_exists parking, and Notes
# ---------------------------------------------------------------------------

@patch("nextseek_api.services.context_catalog._sample_type_rows")
def test_mode_update_renders_only_dseq_with_mapped_percent(rows, tmp_path, monkeypatch):
    rows.return_value = [_D_SEQ_ROW]
    monkeypatch.setattr(
        "nextseek_api.services.reingest_lookups.notes_for_uids",
        lambda uids: {u: "" for u in uids})
    manifest_id = _save_manifest(
        tmp_path, monkeypatch,
        metrics={"star-uniquely_mapped_percent": 91.4})

    result = _dispatch("build-upload-xlsx", {"manifest_id": manifest_id, "mode": "update"},
                        outputs_dir=str(tmp_path))

    assert set(result["saved_files"]) == {"reingest_D_SEQ_update"}
    assert "A.ALN" not in result["qa"]
    assert "A.GEX" not in result["qa"]

    from nextseek_api.batch_upload.convert import parse_traditional_file
    batch = parse_traditional_file(result["saved_files"]["reingest_D_SEQ_update"])
    assert len(batch.rows) == 1
    row = batch.rows[0]
    assert row.UID == "D.SEQ-EXAMPLE-1"
    meta = _meta(row)
    assert meta["MappedPercent"] == 91.4
    assert "Parent" not in meta  # update mode: no Parent column


@patch("nextseek_api.services.context_catalog._sample_type_rows")
def test_an_attribute_not_on_the_schema_is_parked_in_notes_not_a_column(rows, tmp_path, monkeypatch):
    """rRNAPercent maps cleanly (a committed qc_attributes rule) but is not
    declared on _D_SEQ_ROW's schema -- design section 8's "attributes that do
    not exist": parked into Notes, never invented as a real column, and a
    needs_definition proposal is queued for it."""
    rows.return_value = [_D_SEQ_ROW]
    monkeypatch.setattr(
        "nextseek_api.services.reingest_lookups.notes_for_uids",
        lambda uids: {u: "" for u in uids})
    manifest_id = _save_manifest(
        tmp_path, monkeypatch,
        metrics={"star-uniquely_mapped_percent": 91.4,
                 "custom_content_biotype_counts-percent_rRNA": 2.1})

    result = _dispatch("build-upload-xlsx", {"manifest_id": manifest_id, "mode": "update"},
                        outputs_dir=str(tmp_path))

    from nextseek_api.batch_upload.convert import parse_traditional_file
    batch = parse_traditional_file(result["saved_files"]["reingest_D_SEQ_update"])
    row = batch.rows[0]
    meta = _meta(row)
    assert "rRNAPercent" not in meta
    assert meta["MappedPercent"] == 91.4
    notes = meta.get("Notes", "")
    assert "rRNAPercent=2.1" in notes
    assert "[nfcore-reingest" in notes

    needs_def = [p for p in result["proposals"] if p.get("status") == "needs_definition"]
    assert len(needs_def) == 1
    assert needs_def[0]["proposed_attribute"] == "rRNAPercent"
    assert needs_def[0]["proposed_target"] == "D.SEQ"
    assert needs_def[0]["pipeline"] == "nf-core/rnaseq"
    assert needs_def[0]["manifest_digest"] == manifest_id

    # QA must soft-flag the parked value (ATTRIBUTE_NOT_DEFINED), never come
    # back CLEAN on a workbook carrying one -- Plan 3's own "Done when" bar.
    assert result["qa"]["D.SEQ"]["disposition"] == "SOFT_FLAG"
    assert any("attribute_not_defined" in h for h in result["qa"]["D.SEQ"]["soft"])


@patch("nextseek_api.services.context_catalog._sample_type_rows")
def test_a_notes_fetch_failure_drops_the_parked_value_rather_than_guessing(rows, tmp_path, monkeypatch):
    """notes_for_uids omits a UID whose fetch failed (its own documented
    contract). Parking must honour that and write no Notes at all for that
    row rather than risk clobbering text it never read."""
    rows.return_value = [_D_SEQ_ROW]
    monkeypatch.setattr(
        "nextseek_api.services.reingest_lookups.notes_for_uids", lambda uids: {})
    manifest_id = _save_manifest(
        tmp_path, monkeypatch,
        metrics={"custom_content_biotype_counts-percent_rRNA": 2.1})

    result = _dispatch("build-upload-xlsx", {"manifest_id": manifest_id, "mode": "update"},
                        outputs_dir=str(tmp_path))

    from nextseek_api.batch_upload.convert import parse_traditional_file
    batch = parse_traditional_file(result["saved_files"]["reingest_D_SEQ_update"])
    row = batch.rows[0]
    assert "Notes" not in _meta(row)
    # The proposal is still raised even though the Notes write did not
    # happen -- the schema gap is real regardless of whether this run could
    # safely record it in this sample's Notes.
    assert any(p.get("proposed_attribute") == "rRNAPercent" for p in result["proposals"])


def test_attribute_exists_outage_propagates_rather_than_defaulting_false(tmp_path, monkeypatch):
    """An entirely empty catalog is an outage signal (proposals.attribute_exists'
    own contract), not "nothing is defined". The op must let this raise, never
    catch it and default to False -- that would fabricate a schema gap and
    park a good value in Notes."""
    manifest_id = _save_manifest(
        tmp_path, monkeypatch,
        metrics={"star-uniquely_mapped_percent": 91.4})

    with patch("nextseek_api.services.context_catalog._sample_type_rows",
               side_effect=RuntimeError("sample_types_context table unreachable")):
        with pytest.raises(RuntimeError):
            _dispatch("build-upload-xlsx", {"manifest_id": manifest_id, "mode": "update"},
                      outputs_dir=str(tmp_path))


# ---------------------------------------------------------------------------
# Provenance sheet wiring (brief Step 3's third requirement)
# ---------------------------------------------------------------------------

@patch("nextseek_api.services.context_catalog._sample_type_rows")
def test_provenance_sheet_is_written_when_an_attribute_has_a_non_map_origin(rows, tmp_path, monkeypatch):
    rows.return_value = [_D_SEQ_ROW]
    monkeypatch.setattr(
        "nextseek_api.services.reingest_lookups.notes_for_uids",
        lambda uids: {u: "" for u in uids})
    manifest_id = _save_manifest(
        tmp_path, monkeypatch,
        metrics={"star-uniquely_mapped_percent": 91.4,
                 "custom_content_biotype_counts-percent_rRNA": 2.1})

    result = _dispatch("build-upload-xlsx", {"manifest_id": manifest_id, "mode": "update"},
                        outputs_dir=str(tmp_path))

    wb = openpyxl.load_workbook(result["saved_files"]["reingest_D_SEQ_update"])
    assert "Provenance" in wb.sheetnames
    sheet = wb["Provenance"]
    header = [c.value for c in sheet[1]]
    assert header == ["UID", "Attribute", "Value", "Origin", "Raw key", "Source file"]
    rows_by_attr = {r[1].value: r for r in sheet.iter_rows(min_row=2)}
    assert rows_by_attr["MappedPercent"][3].value == "map"
    assert rows_by_attr["rRNAPercent"][3].value == "parked"
    assert rows_by_attr["rRNAPercent"][4].value == "custom_content_biotype_counts-percent_rRNA"


def test_no_provenance_sheet_wiring_means_no_sheet_would_appear(tmp_path, monkeypatch):
    """Negative control for the wiring check above: render_upload_workbook
    only emits Provenance when it is actually PASSED a provenance list (see
    its own docstring) -- confirms the assertion above is exercising real
    wiring, not something render_upload_workbook does unconditionally."""
    from NessieAI.ns.upload_workbook import render_upload_workbook

    out = tmp_path / "no_provenance.xlsx"
    render_upload_workbook(
        "A.ALN", [{"json_metadata": {"Parent": "D.SEQ-1", "Type": "BAM"}, "assay_ids": []}],
        str(out))
    wb = openpyxl.load_workbook(out)
    assert "Provenance" not in wb.sheetnames


# ---------------------------------------------------------------------------
# Legacy rows path stays untouched (brief: "Keep --rows working")
# ---------------------------------------------------------------------------

def test_legacy_rows_path_still_works_without_a_manifest_id(tmp_path):
    import json as _json

    rows = _json.dumps([{
        "SampleType": "A.SCXP",
        "json_metadata": {"Parent": "D.SEQ-1", "Scientist": "Marie Floryan"},
        "assay_ids": [12],
    }])
    result = _dispatch("build-upload-xlsx", {"rows": rows, "existing_parent_uids": "D.SEQ-1"},
                        outputs_dir=str(tmp_path))
    assert set(result["saved_files"]) == {"reingest_A_SCXP"}
    assert "reply" not in result  # legacy envelope shape, unchanged
    assert "proposals" not in result
