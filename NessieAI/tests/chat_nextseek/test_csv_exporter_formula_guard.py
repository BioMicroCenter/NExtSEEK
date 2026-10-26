"""SEC-H03: the chat CSV exporter prefixes formula-looking strings, not numbers."""
import csv

from chat_nextseek.reports.exporters.csv import _write_csv_rows


def test_formula_prefixed_numbers_and_plain_untouched(tmp_path):
    p = tmp_path / "o.csv"
    _write_csv_rows(p, [{"a": '=HYPERLINK("x","y")', "b": "-80", "c": "+1.5e3", "d": "plain", "e": "@SUM(1)"}],
                    ["a", "b", "c", "d", "e"])
    row = next(csv.DictReader(p.open()))
    assert row == {"a": "'=HYPERLINK(\"x\",\"y\")", "b": "-80", "c": "+1.5e3", "d": "plain", "e": "'@SUM(1)"}


EVIL = '=HYPERLINK("http://x","y")'


def test_geo_cells_are_text():
    from openpyxl import Workbook
    from chat_nextseek.reports.exporters.geo_xlsx import _write_cell, _write_list_down

    ws = Workbook().active
    _write_cell(ws, 1, 1, EVIL)
    _write_list_down(ws, 2, 1, [EVIL])
    for ref in ("A1", "A2"):
        assert ws[ref].data_type == "s" and ws[ref].value == EVIL


def test_sra_rows_are_text():
    from openpyxl import Workbook
    from chat_nextseek.reports.exporters.sra_xlsx import _write_template_rows

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "title"
    _write_template_rows(wb, sheet_name="S", header_row=1, template_row=2, rows=[{"title": EVIL}])
    assert ws["A2"].data_type == "s" and ws["A2"].value == EVIL
