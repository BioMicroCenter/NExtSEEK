# NessieAI/ns/reingest/workbook_diff.py
"""Is a rebuilt workbook the reviewed one plus Assay rows for new assays? (spec 3.5 step 3)

The upload creates a missing study assay, rebuilds, and uploads the rebuilt
workbook in place of the reviewed one only when nothing else changed. The
reviewed Assay sheet shows each assay to create as its planned label
(assay_plan.planned_assay_label); the rebuild must hold that assay's created id
in its place. Pure: openpyxl only.
"""
from __future__ import annotations

import re
from collections import Counter

ASSAY_SHEET = "Assay"
ASSAY_ID_COLUMN = 2  # SampleType, AssayType, Assay, Direction
# The Provenance sheet's "Answered by" cell is "<curator> on <build date>"; a
# rebuild stamps its own date, so a build reviewed one day and confirmed the
# next would otherwise always differ. Only that trailing date is ignored.
PROVENANCE_SHEET = "Provenance"
ANSWERED_BY = "Answered by"
_ANSWER_DATE = re.compile(r" on \d{4}-\d{2}-\d{2}$")


def _rows(book, name) -> list[tuple]:
    rows = [tuple(row) for row in book[name].iter_rows(values_only=True)]
    if name != PROVENANCE_SHEET or not rows or ANSWERED_BY not in rows[0]:
        return rows
    column = rows[0].index(ANSWERED_BY)
    return [rows[0]] + [
        row[:column] + (_ANSWER_DATE.sub("", row[column]),) + row[column + 1:]
        if len(row) > column and isinstance(row[column], str) else row
        for row in rows[1:]]


def _with_created_ids(row: tuple, created: dict) -> tuple:
    """A reviewed Assay row with its planned label replaced by the created id."""
    if len(row) <= ASSAY_ID_COLUMN or row[ASSAY_ID_COLUMN] not in created:
        return row
    return row[:ASSAY_ID_COLUMN] + (created[row[ASSAY_ID_COLUMN]],) + row[ASSAY_ID_COLUMN + 1:]


def assay_only_changes(reviewed_path: str, rebuilt_path: str, created: dict) -> list[str]:
    """``created`` maps each planned label the upload acted on to the assay id it created."""
    import openpyxl

    created = {str(label): int(i) for label, i in created.items()}
    allowed = set(created.values())
    reviewed = openpyxl.load_workbook(reviewed_path, read_only=True, data_only=True)
    rebuilt = openpyxl.load_workbook(rebuilt_path, read_only=True, data_only=True)
    try:
        reasons: list[str] = []
        if reviewed.sheetnames != rebuilt.sheetnames:
            return [f"the sheets changed: {reviewed.sheetnames} became {rebuilt.sheetnames}"]
        for name in reviewed.sheetnames:
            before, after = _rows(reviewed, name), _rows(rebuilt, name)
            if name != ASSAY_SHEET:
                if before != after:
                    reasons.append(f"{name}: differs from the reviewed workbook")
                continue
            before = before[:1] + [_with_created_ids(row, created) for row in before[1:]]
            missing = Counter(before) - Counter(after)
            added = Counter(after) - Counter(before)
            if missing:
                reasons.append(f"{ASSAY_SHEET}: a reviewed row is gone")
            stray = [row for row in added
                     if len(row) <= ASSAY_ID_COLUMN or row[ASSAY_ID_COLUMN] not in allowed]
            if stray:
                reasons.append(f"{ASSAY_SHEET}: rows for assays this upload did not create")
        return reasons
    finally:
        reviewed.close()
        rebuilt.close()
