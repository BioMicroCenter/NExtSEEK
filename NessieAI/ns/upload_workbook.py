"""Render a NExtSEEK 4-sheet upload workbook from CC-composed reingest rows.

One workbook == one SampleType (batch_upload/convert.py rejects INSTRUCTIONS that
declare more than one SampleType). The output parses cleanly back through
``nextseek_api.batch_upload.convert.parse_traditional_file`` — that round-trip is
the correctness oracle (see tests/test_upload_workbook.py and
tests/test_upload_workbook_modes.py).

Row shape (one per output sample):
    {"json_metadata": {<attr>: <value>, ...}, "assay_ids": [int, ...]}

Two modes:
- ``mode="new"`` (default) — every sample is [NEW]; no UID column, the server
  mints the UID on upload. ``json_metadata`` carries Parent / Scientist /
  File_PrimaryData and any other attributes CC derived.
- ``mode="update"`` — a QC-backfill workbook: every row targets an existing
  sample and must carry that sample's UID in ``json_metadata["UID"]``. The UID
  is rendered as the first Samples column but is deliberately NOT declared in
  Instructions — convert.py's ``_prepare_traditional_row`` intercepts a UID
  column before the Instructions lookup, so declaring it there would be
  redundant (and Instructions still filters every other column). This mode is
  meant to be uploaded with ``update_existing: true``, a request-level flag,
  not per-row.
"""
from __future__ import annotations

import json
from typing import Any

import openpyxl

MODE_NEW = "new"
MODE_UPDATE = "update"


def render_upload_workbook(
    sample_type: str,
    rows: list[dict[str, Any]],
    out_path: str,
    *,
    mode: str = MODE_NEW,
    provenance: list[dict[str, Any]] | None = None,
) -> None:
    """Write a NExtSEEK upload workbook for one ``sample_type`` to ``out_path``.

    Emits all four sheets (Instructions / Samples / Assay / Ontology); a missing
    sheet would make convert.py silently fall back to the flat parser. When
    ``provenance`` is given, a fifth ``Provenance`` sheet is also written — the
    traditional parser reads only the four named sheets, so it is inert on
    upload.
    """
    if not rows:
        raise ValueError("render_upload_workbook: no rows to render")
    if mode not in (MODE_NEW, MODE_UPDATE):
        raise ValueError(f"render_upload_workbook: unknown mode {mode!r}")

    # Field set = union of json_metadata keys, first-seen order. UID is handled
    # separately: convert.py intercepts it before the Instructions lookup, so
    # declaring it there would be redundant.
    fields: list[str] = []
    for row in rows:
        for key in (row.get("json_metadata") or {}):
            if key != "UID" and key not in fields:
                fields.append(key)
    if not fields:
        raise ValueError("render_upload_workbook: rows carry no json_metadata")

    if mode == MODE_UPDATE:
        for index, row in enumerate(rows):
            if not str((row.get("json_metadata") or {}).get("UID") or "").strip():
                raise ValueError(f"render_upload_workbook: row {index} has no UID "
                                 f"(update mode targets existing samples)")

    wb = openpyxl.Workbook()

    # INSTRUCTIONS: one row per Field. Database Field MUST be "SampleType::Attr"
    # (the "::" is required by InstructionRow validation). Field == attribute name.
    wi = wb.active
    wi.title = "Instructions"
    wi.append(["Field", "Database Field", "Field Type", "Ontology"])
    for field in fields:
        wi.append([field, f"{sample_type}::{field}", _field_type(field, rows), None])

    # SAMPLES: one column per Field (headers must equal the Field strings exactly —
    # convert.py matches trim-only, case-sensitive). In update mode, UID is the
    # first column and carries the existing sample's identifier; in new mode
    # there is no UID column at all, so it stays blank and the server mints one.
    ws = wb.create_sheet("Samples")
    header = (["UID"] if mode == MODE_UPDATE else []) + list(fields)
    ws.append(header)
    for row in rows:
        meta = row.get("json_metadata") or {}
        prefix = [meta.get("UID")] if mode == MODE_UPDATE else []
        ws.append(prefix + [_cell(meta.get(field)) for field in fields])

    # ASSAY: one row per distinct assay_id, applied to this sample type (Direction 1).
    wa = wb.create_sheet("Assay")
    wa.append(["SampleType", "AssayType", "Assay", "Direction"])
    seen: set[int] = set()
    for row in rows:
        for assay_id in (row.get("assay_ids") or []):
            if assay_id not in seen:
                seen.add(assay_id)
                wa.append([sample_type, None, assay_id, 1])

    # ONTOLOGY: present but empty — reingest fields are free-text, no controlled vocab.
    wo = wb.create_sheet("Ontology")
    wo.append([])

    # PROVENANCE: optional fifth sheet, ignored by the traditional parser (it
    # reads only the four sheets above by name), so it is inert on upload.
    if provenance:
        wp = wb.create_sheet("Provenance")
        wp.append(["UID", "Attribute", "Value", "Origin", "Raw key", "Source file"])
        for entry in provenance:
            wp.append([entry.get("uid", ""), entry.get("attribute", ""),
                       _cell(entry.get("value")), entry.get("origin", ""),
                       entry.get("raw_key", ""), entry.get("source_file", "")])

    wb.save(out_path)


def _field_type(field: str, rows: list[dict[str, Any]]) -> str:
    """"Number" when every non-blank value for ``field`` is numeric, else "Text".

    Never "Controlled Ontology": that triggers ontology.py validation against an
    Ontology sheet reingest leaves empty.
    """
    seen = False
    for row in rows:
        value = (row.get("json_metadata") or {}).get(field)
        if value is None or value == "":
            continue
        seen = True
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            return "Text"
    return "Number" if seen else "Text"


def _cell(value: Any) -> Any:
    """Serialize a metadata value into an xlsx cell (nested dict/list -> compact JSON)."""
    if isinstance(value, (dict, list)):
        return json.dumps(value, separators=(",", ":"))
    return value
