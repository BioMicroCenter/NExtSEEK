"""The supplementary zip → evidence items, and gaps for what could not be read.

Workbooks and csv/tsv give one item per non-empty cell, a workbook cell rendered as a curator
sees it (1, not 1.0; 2024-02-03, not a datetime; 007 under a 000 format), so a sample name
matches verbatim. PDF and txt give paragraphs. Anything else, anything over the size cap, and
anything the paper lists but the zip lacks becomes a gap (spec §10: a missing supplement never
stops a run). An unreadable archive, or a readable file with no text, is a gap too.
Decompression is bounded three ways: each member by its declared size, each workbook (itself a
zip) by the declared sizes of its parts, and the archive as a whole by one budget. zipfile stops
at a member's declared size and fails its CRC, so a header that lies cannot unpack past it.
Europe PMC's zip also carries the article's figure images; an image the paper does not list
as a supplement is skipped, not a gap. A listed name matches a member by file name, URL-decoded
and case-folded. Evidence IDs stay unique: a stem already used in the archive is replaced by
the member's path, then by the path with #n.
"""
from __future__ import annotations

import csv
import io
import os
import re
import zipfile
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal
from urllib.parse import unquote

from NessieAI.ns.paper.models import EvidenceItem, Gap, SupplementRef
from NessieAI.ns.paper.pdftext import chunk_paragraphs, pdf_to_text

MAX_MEMBER_BYTES = 50 * 1024 * 1024
MAX_WORKBOOK_UNPACKED_BYTES = 256 * 1024 * 1024  # a workbook's parts, as declared inside it
MAX_ARCHIVE_UNPACKED_BYTES = 1024 * 1024 * 1024  # every member, and every workbook's parts, together
MAX_ITEMS_PER_FILE = 50_000
_DELIMITED = {".csv": ",", ".tsv": "\t"}
_FIGURE_EXTS = frozenset({".gif", ".jpg", ".jpeg", ".png", ".tif", ".tiff"})


def supplement_items(zip_bytes: bytes, expected: list[SupplementRef]) -> tuple[list[EvidenceItem], list[Gap]]:
    try:
        zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError) as exc:
        return [], [Gap(kind="supplement_unparsed", detail=f"supplementary archive: {type(exc).__name__}")] + [
            Gap(kind="supplement_unavailable", detail=s.name) for s in expected]
    items: list[EvidenceItem] = []
    gaps: list[Gap] = []
    listed = {_match_key(s.name) for s in expected}
    found: set[str] = set()
    used: set[str] = set()
    unpacked = 0
    with zf:
        for info in zf.infolist():
            base = info.filename.rsplit("/", 1)[-1]
            if info.is_dir() or not base or base.startswith("."):
                continue
            key = _match_key(base)
            if key not in listed and os.path.splitext(key)[1] in _FIGURE_EXTS:
                continue  # the article's own figure images ride in the same zip; they are read as captions
            found.add(key)
            name = _claim_stem(info.filename, base, used)
            if info.file_size > MAX_MEMBER_BYTES:
                gaps.append(Gap(kind="supplement_unparsed", detail=f"{name}: over {MAX_MEMBER_BYTES // 2**20} MB"))
                continue
            if unpacked + info.file_size > MAX_ARCHIVE_UNPACKED_BYTES:
                gaps.append(_over_budget(name))
                continue
            ext = os.path.splitext(base)[1].lower()
            try:
                data = zf.read(info)
                unpacked += len(data)
                if ext in (".xlsx", ".xlsm"):
                    parts = _workbook_unpacked(data)
                    if parts > MAX_WORKBOOK_UNPACKED_BYTES:
                        gaps.append(Gap(kind="supplement_unparsed", detail=(
                            f"{name}: unpacked, over {MAX_WORKBOOK_UNPACKED_BYTES // 2**20} MB")))
                        continue
                    if unpacked + parts > MAX_ARCHIVE_UNPACKED_BYTES:
                        gaps.append(_over_budget(name))
                        continue
                    unpacked += parts
                    new = _workbook(name, data)
                elif ext in _DELIMITED:
                    new = _delimited(name, data, _DELIMITED[ext])
                elif ext == ".pdf":
                    new = _prose(name, pdf_to_text(data))
                elif ext == ".txt":
                    new = _prose(name, data.decode("utf-8", errors="replace"))
                else:
                    gaps.append(Gap(kind="supplement_unparsed", detail=f"{name}: {ext or 'no extension'} is not read"))
                    continue
            except Exception as exc:  # a corrupt member is a gap, never a failed run
                gaps.append(Gap(kind="supplement_unparsed", detail=f"{name}: {type(exc).__name__}"))
                continue
            if not new:
                gaps.append(Gap(kind="supplement_unparsed", detail=f"{name}: no readable text"))
                continue
            if len(new) > MAX_ITEMS_PER_FILE:
                gaps.append(Gap(kind="supplement_unparsed", detail=f"{name}: only the first {MAX_ITEMS_PER_FILE} cells read"))
                new = new[:MAX_ITEMS_PER_FILE]
            items += new
    missing: dict[str, str] = {}
    for s in expected:
        if _match_key(s.name) not in found:
            missing.setdefault(_match_key(s.name), s.name)
    gaps += [Gap(kind="supplement_unavailable", detail=name) for name in missing.values()]
    return items, gaps


def _match_key(name: str) -> str:
    """How a listed supplement is matched to a zip member: file name only, URL-decoded, case-folded."""
    return unquote(name.rsplit("/", 1)[-1]).strip().casefold()


def _over_budget(name: str) -> Gap:
    return Gap(kind="supplement_unparsed",
               detail=f"{name}: the archive's {MAX_ARCHIVE_UNPACKED_BYTES // 2**20} MB unpacked budget is spent")


def _workbook_unpacked(data: bytes) -> int:
    """The declared unpacked size of a workbook's parts; openpyxl's zipfile reads no more than this."""
    with zipfile.ZipFile(io.BytesIO(data)) as inner:
        return sum(i.file_size for i in inner.infolist())


def _claim_stem(path: str, base: str, used: set[str]) -> str:
    """The ID stem for one member: its basename, else its zip path, else the path with #n. Recorded in used."""
    if base not in used:
        stem = base
    elif path not in used:
        stem = path
    else:
        n = 2
        while f"{path}#{n}" in used:
            n += 1
        stem = f"{path}#{n}"
    used.add(stem)
    return stem


def _workbook(name: str, data: bytes) -> list[EvidenceItem]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[EvidenceItem] = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    text = "" if cell.value is None else cell_text(cell.value, getattr(cell, "number_format", None))
                    if not text:
                        continue
                    out.append(EvidenceItem(id=f"E:si:{name}!{ws.title}!{cell.coordinate}", text=text,
                                            source=f"SI {name}, sheet {ws.title}, cell {cell.coordinate}"))
                    if len(out) > MAX_ITEMS_PER_FILE:
                        return out
    finally:
        wb.close()
    return out


_PADDED = re.compile(r"0+")
_FIXED = re.compile(r"0(?:\.(0+))?(%?)")


def cell_text(value, number_format: str | None = None) -> str:
    """A workbook value as a curator sees it in the sheet, stripped; dates in ISO (the locale is unknown)."""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time():
            return value.date().isoformat()
        return f"{value.date().isoformat()} {_clock(value.time())}"
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return _clock(value)
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, (int, float)):
        return _number(value, (number_format or "General").split(";")[0].strip())
    return str(value).strip()


def _clock(t: time) -> str:
    return t.isoformat(timespec="minutes" if not (t.second or t.microsecond) else "seconds")


def _number(value: int | float, fmt: str) -> str:
    """General as Excel shows it (up to 15 significant digits), plus the plain 000, 0.00 and 0.0% formats."""
    exact = Decimal(repr(value)) if isinstance(value, float) else Decimal(value)
    if not exact.is_finite():
        return str(value)
    if _PADDED.fullmatch(fmt):
        return f"{int(exact.quantize(Decimal(1), ROUND_HALF_UP)):0{len(fmt)}d}"
    fixed = _FIXED.fullmatch(fmt)
    if fixed:
        places, percent = len(fixed.group(1) or ""), fixed.group(2)
        shown = (exact * 100 if percent else exact).quantize(Decimal(1).scaleb(-places), ROUND_HALF_UP)
        return f"{shown:f}{percent}"
    if isinstance(value, int):
        return str(value)
    if value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return f"{value:.15g}".replace("e", "E")


def _delimited(name: str, data: bytes, sep: str) -> list[EvidenceItem]:
    from openpyxl.utils import get_column_letter

    out: list[EvidenceItem] = []
    reader = csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace")), delimiter=sep)
    for r, row in enumerate(reader, start=1):
        for c, value in enumerate(row, start=1):
            if value.strip():
                coord = f"{get_column_letter(c)}{r}"
                out.append(EvidenceItem(id=f"E:si:{name}!{coord}", text=value.strip(), source=f"SI {name}, cell {coord}"))
                if len(out) > MAX_ITEMS_PER_FILE:
                    return out
    return out


def _prose(name: str, text: str) -> list[EvidenceItem]:
    return [EvidenceItem(id=f"E:si:{name}!p{i}", text=chunk, source=f"SI {name}, paragraph {i}")
            for i, chunk in enumerate(chunk_paragraphs(text.splitlines()), start=1)]
