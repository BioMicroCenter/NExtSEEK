"""``POST /nextseek_api/samples/retrieve/`` and its alias ``admin/samples/retrieve/``: the sample download API.

Every sample-download control in the UI, the Nessie reporter and the pipeline agent read through this endpoint, so
these tests pin the whole contract, end to end through the URL conf:

- the request: ``identifiers`` (or the three legacy names), numeric SEEK ids, ``include_tree``, ``output_format``;
- the JSON body, key order included, and the workbook response;
- scope: a superuser is unscoped, anyone else sees only their projects, and a sample outside them answers exactly as
  an unknown one (#74);
- lineage: every ancestor and descendant, not siblings, through an OrphanSample; and what happens when the graph
  lags MySQL or is down: the requested samples are still exported, and ``lineage_complete`` says the lineage is not.

MySQL is real SQL on an in-memory sqlite3 database shaped like SEEK's tables, handed to the module (and to
``resolve_scope``) through a cursor that takes Django's ``%s`` placeholders. Neo4j is a fake graph behind
``sample_retrieve._neo4j_run`` that answers the module's two statements from an adjacency list, so the Cypher itself
is not exercised here (the parity script runs it against a real graph).
"""

import base64
import json
import sqlite3
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from nextseek_api.services import sample_retrieve as sr

NEW = "/nextseek_api/samples/retrieve/"
OLD = "/nextseek_api/admin/samples/retrieve/"
PATHS = pytest.mark.parametrize("path", [NEW, OLD], ids=["samples-retrieve", "admin-alias"])

# The default database only: the request cycle's middleware reads it. SEEK's tables are the sqlite3 fixture below.
pytestmark = pytest.mark.django_db

# MySQL: id -> (uuid, sample_type_id, json_metadata, project ids)
ROWS = {
    1: ("NHP-1", 41, '{"UID": "NHP-1", "Name": "monkey"}', {2}),
    2: ("TIS-1", 7, '{"UID": "TIS-1", "Parent": "NHP-1"}', {2}),
    3: ("TIS-2", 7, '{"UID": "TIS-2", "Parent": "NHP-1"}', {2}),
    4: ("RNA-1", 9, '{"UID": "RNA-1", "Parent": "TIS-1"}', {2, 3}),  # in two of the member's projects
    5: ("D.SEQ-1", 11, '{"UID": "D.SEQ-1", "Parent": "RNA-1"}', {2}),
    6: ("TIS-FOR-1", 7, '{"UID": "TIS-FOR-1", "Parent": "NHP-1"}', {4}),  # another lab's sample
    7: ("DNA-OFF-1", 12, '{"UID": "DNA-OFF-1", "Parent": "TIS-1"}', {4}),  # the member's relative, not theirs
    8: ("TIS-NEW-1", 7, '{"UID": "TIS-NEW-1"}', {2}),  # uploaded; the graph sync has not reached it
    9: ("MUS-1", 3, "", {2}),
    10: ("MUS-2", 3, "not json", {2}),
    11: ("SLD-1", 5, '{"UID": "SLD-1"}', {2}),  # reachable from NHP-1 only through the orphan
}

# The graph: uuid -> (id, live Sample?) and DERIVED_FROM child -> parents. The orphan is a graph-only v1.0 node.
NODES = {
    "NHP-1": (1, True), "TIS-1": (2, True), "TIS-2": (3, True), "RNA-1": (4, True), "D.SEQ-1": (5, True),
    "TIS-FOR-1": (6, True), "DNA-OFF-1": (7, True), "MUS-1": (9, True), "MUS-2": (10, True),
    "ORPH-1": (999, False), "SLD-1": (11, True),
}
PARENTS = {
    "TIS-1": ["NHP-1"], "TIS-2": ["NHP-1"], "RNA-1": ["TIS-1"], "D.SEQ-1": ["RNA-1"], "TIS-FOR-1": ["NHP-1"],
    "DNA-OFF-1": ["TIS-1"], "ORPH-1": ["NHP-1"], "SLD-1": ["ORPH-1"],
}

MEMBER, SUPER, NOBODY = "member", "admin", "stranger"


class Graph:
    """The two statements of sample_retrieve, answered from NODES and PARENTS."""

    def __init__(self):
        self.nodes, self.parents = dict(NODES), {k: list(v) for k, v in PARENTS.items()}
        self.twins = {}  # uuid -> (id, live) of a second node carrying the same uuid
        self.calls = []  # (statement name, sorted uuids)
        self.down = False

    def _reach(self, start, step):
        out, stack = set(), [start]
        while stack:
            for nxt in step(stack.pop()):
                if nxt not in out:
                    out.add(nxt)
                    stack.append(nxt)
        return out

    def run(self, cypher, uuids):
        name = "resolve" if cypher is sr.RESOLVE_CYPHER else "lineage"
        self.calls.append((name, sorted(uuids)))
        if self.down:
            raise sr.GraphUnavailable("connection refused")
        live = [u for u in uuids if u in self.nodes and self.nodes[u][1]]
        if name == "resolve":
            return [{"uuid": u, "id": self.nodes[u][0]} for u in live]
        children = lambda n: [c for c, ps in self.parents.items() if n in ps]  # noqa: E731
        found = set(live)
        for u in live:
            found |= self._reach(u, lambda n: self.parents.get(n, []))
            found |= self._reach(u, children)
        twins = lambda u: [i for key, (i, live) in self.twins.items() if key == u and live]  # noqa: E731
        return [{"uuid": u, "id": self.nodes[u][0], "live": self.nodes[u][1],
                 "twin_ids": ([self.nodes[u][0]] if self.nodes[u][1] else []) + twins(u)} for u in sorted(found)]

    def walked(self):
        return [u for name, us in self.calls if name == "lineage" for u in us]


@pytest.fixture
def graph(monkeypatch):
    g = Graph()
    monkeypatch.setattr(sr, "_neo4j_run", lambda cypher, uuids: g.run(cypher, uuids))
    # The workbook's README tree asks the graph for hop assays on its own (bounded, fail-soft); not here.
    monkeypatch.setattr("nextseek_api.services.sample_workbook.load_derivation_hops", lambda uuids: [])
    return g


class _Cursor:
    """A sqlite3 cursor that takes Django's ``%s`` placeholders, as a Django cursor does."""

    def __init__(self, conn):
        self._c = conn.cursor()

    def execute(self, sql, params=()):
        self._c.execute(sql.replace("%s", "?"), list(params or ()))

    def fetchall(self):
        return self._c.fetchall()

    def fetchone(self):
        return self._c.fetchone()

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self._c.close()
        return False


class _Connection:
    def __init__(self, conn):
        self.conn = conn

    def cursor(self):
        return _Cursor(self.conn)


@pytest.fixture
def seek(monkeypatch):
    """SEEK's tables, as much of them as the endpoint and resolve_scope read, shaped as in SEEK's MySQL."""
    conn = sqlite3.connect(":memory:")
    c = conn.cursor()
    c.execute("CREATE TABLE samples (id INTEGER PRIMARY KEY, sample_type_id INTEGER, uuid TEXT, json_metadata TEXT)")
    c.execute("CREATE TABLE projects_samples (project_id INTEGER, sample_id INTEGER)")
    c.execute("CREATE TABLE users (login TEXT, person_id INTEGER)")
    c.execute("CREATE TABLE work_groups (id INTEGER, project_id INTEGER)")
    c.execute("CREATE TABLE group_memberships (person_id INTEGER, work_group_id INTEGER)")
    for sid, (uuid, st, meta, projects) in ROWS.items():
        c.execute("INSERT INTO samples VALUES (?, ?, ?, ?)", [sid, st, uuid, meta])
        for p in projects:
            c.execute("INSERT INTO projects_samples VALUES (?, ?)", [p, sid])
    # member belongs to projects 2 and 3; stranger has a SEEK person but no project.
    c.execute("INSERT INTO users VALUES ('member', 100), ('stranger', 101)")
    c.execute("INSERT INTO work_groups VALUES (20, 2), (30, 3), (40, 4)")
    c.execute("INSERT INTO group_memberships VALUES (100, 20), (100, 30)")
    conn.commit()
    wrapped = _Connection(conn)
    monkeypatch.setattr(sr, "_cursor", wrapped.cursor)
    monkeypatch.setattr("nextseek_api.graph_search.scope.connections", {"seek": wrapped})
    yield wrapped
    conn.close()


def _client(login=MEMBER, *, staff=True):
    # A stand-in user: saving a real one fires Mezzanine's profile hook, which the database router refuses here.
    # is_staff is True by default because the SEEK login sets it on everyone; it must not widen scope.
    user = SimpleNamespace(username=login, is_authenticated=True, is_active=True, is_anonymous=False,
                           is_staff=staff, is_superuser=login == SUPER, pk=1, id=1)
    client = APIClient()
    client.force_authenticate(user)
    client.credentials(HTTP_AUTHORIZATION="Basic " + base64.b64encode(f"{login}:pw".encode()).decode())
    return client


def _post(body, *, login=MEMBER, path=NEW, fmt="json"):
    return _client(login).post(path, data=json.dumps(body) if fmt == "json" else body,
                               content_type="application/json" if fmt == "json" else None)


def _uuids(resp):
    return sorted(s["uuid"] for g in resp.json()["data"] for s in g["samples"])


# --------------------------------------------------------------------------- the contract


@PATHS
def test_json_body_shape_and_key_order(seek, graph, path):
    resp = _post({"identifiers": ["TIS-2"]}, login=SUPER, path=path)

    assert resp.status_code == 200
    body = resp.json()
    assert list(body) == ["total_samples", "total_sample_types", "total_children", "failed_uids", "data",
                          "lineage_complete"]
    assert body["lineage_complete"] is True and resp["X-NExtSEEK-Lineage-Complete"] == "true"
    nhp = next(g for g in body["data"] if g["sample_type"] == "NHP")
    assert list(nhp) == ["sample_type", "n_samples", "samples"]
    assert nhp["samples"] == [{"id": "1", "uuid": "NHP-1", "sample_type_id": 41,
                               "metadata": {"UID": "NHP-1", "Name": "monkey"}}]


def test_both_paths_answer_identically(seek, graph):
    for body in ({"identifiers": ["TIS-1"]}, {"identifiers": ["TIS-FOR-1"]}, {"identifiers": ["TIS-1"], "include_tree": False}):
        new, old = _post(body, path=NEW), _post(body, path=OLD)
        assert (new.status_code, new.content) == (old.status_code, old.content)


def test_get_is_not_allowed_on_the_new_path(seek, graph):
    assert _client().get(NEW).status_code == 405


def test_the_new_path_is_not_the_sample_detail_route():
    from django.urls import resolve

    assert resolve(NEW).func.cls.__name__ == "SampleRetrieveViewSet"
    assert resolve(OLD).func.cls.__name__ == "AdminSampleViewSet"


def test_a_token_caller_is_refused(seek, graph):
    client = _client()
    client.credentials()  # no Basic header, and the session path finds no SEEK login
    from unittest.mock import patch

    with patch.object(sr, "resolve_seek_auth", return_value=(None, None)):
        resp = client.post(NEW, data={"identifiers": ["TIS-1"]}, format="json")
    assert resp.status_code == 401 and resp.json() == {"detail": "Authentication required"}


def test_an_anonymous_caller_is_refused(seek, graph):
    assert APIClient().post(NEW, data={"identifiers": ["TIS-1"]}, format="json").status_code in (401, 403)


@pytest.mark.parametrize("body", [{"identifiers": []}, {"identifiers": ["  ", ""]}, {}, ["TIS-1"]],
                         ids=["empty", "blank", "missing", "not-a-dict"])
def test_nothing_to_retrieve_is_400(seek, graph, body):
    resp = _post(body)
    assert resp.status_code == 400 and resp.json() == {"detail": "identifiers required"}


def test_an_unknown_output_format_is_422(seek, graph):
    resp = _post({"identifiers": ["TIS-1"], "output_format": "xlsx"})
    assert resp.status_code == 422 and resp.json()["detail"] == "Invalid request"


@pytest.mark.parametrize("body", [
    {"retrieval_uids": ["TIS-2"]}, {"uids": ["TIS-2"]}, {"retrieval_uids_text": "TIS-2"},
    {"identifiers": "TIS-2  "},
], ids=["retrieval_uids", "uids", "retrieval_uids_text", "whitespace-string"])
def test_legacy_field_names_and_a_string_still_work(seek, graph, body):
    resp = _post({**body, "include_tree": False})
    assert resp.status_code == 200 and _uuids(resp) == ["TIS-2"]


def test_form_encoded_include_tree_false(seek, graph):
    resp = _client().post(NEW, data={"identifiers": "TIS-2", "include_tree": "false"})
    assert resp.status_code == 200 and _uuids(resp) == ["TIS-2"]
    assert graph.walked() == []


def test_numeric_ids_are_seek_ids_and_unknown_ones_count_as_failed(seek, graph):
    resp = _post({"identifiers": ["3", "424242"], "include_tree": False})
    assert _uuids(resp) == ["TIS-2"] and resp.json()["failed_uids"] == 1


def test_a_failed_numeric_lookup_counts_every_numeric_id_as_failed(seek, graph, monkeypatch):
    real = sr._ids_to_uuids
    calls = []

    def flaky(ids):
        calls.append(ids)
        if len(calls) == 1:
            raise RuntimeError("db down")
        return real(ids)

    monkeypatch.setattr(sr, "_ids_to_uuids", flaky)
    resp = _post({"identifiers": ["3", "TIS-2"], "include_tree": False})
    assert _uuids(resp) == ["TIS-2"] and resp.json()["failed_uids"] == 1


def test_duplicates_are_read_once(seek, graph):
    resp = _post({"identifiers": ["TIS-2", "TIS-2", "3"], "include_tree": False})
    body = resp.json()
    assert _uuids(resp) == ["TIS-2"] and body["total_samples"] == 1 and body["failed_uids"] == 0


def test_metadata_variants(seek, graph):
    resp = _post({"identifiers": ["MUS-1", "MUS-2", "TIS-2"], "include_tree": False})
    meta = {s["uuid"]: s["metadata"] for g in resp.json()["data"] for s in g["samples"]}
    assert meta == {"MUS-1": {}, "MUS-2": {}, "TIS-2": {"UID": "TIS-2", "Parent": "NHP-1"}}


def test_groups_are_uid_prefixes_dotted_ones_included(seek, graph):
    resp = _post({"identifiers": ["D.SEQ-1", "TIS-2", "RNA-1"], "include_tree": False})
    assert [g["sample_type"] for g in resp.json()["data"]] == ["D.SEQ", "RNA", "TIS"]


def test_a_quote_in_a_uid_is_bound_not_interpolated(seek, graph):
    resp = _post({"identifiers": ["TIS-1' OR '1'='1", "x\"); DROP TABLE samples; --"], "include_tree": False})
    assert resp.status_code == 404
    with seek.cursor() as c:
        c.execute("SELECT COUNT(*) FROM samples")
        assert c.fetchone()[0] == len(ROWS)


def test_a_mysql_failure_is_500_with_its_message(seek, graph, monkeypatch):
    def boom(ids, scope):
        raise RuntimeError("Lost connection to MySQL server")

    monkeypatch.setattr(sr, "_hydrate", boom)
    resp = _post({"identifiers": ["TIS-2"]})
    assert resp.status_code == 500 and resp.json() == {"detail": "Lost connection to MySQL server"}


def test_nothing_found_is_404(seek, graph):
    resp = _post({"identifiers": ["TIS-NOPE"]}, login=SUPER)
    assert resp.status_code == 404 and resp.json() == {"detail": "No samples found for provided UIDs"}


# --------------------------------------------------------------------------- lineage


def test_lineage_is_ancestors_and_descendants_not_siblings(seek, graph):
    resp = _post({"identifiers": ["RNA-1"]}, login=SUPER)
    body = resp.json()
    # TIS-2 and TIS-FOR-1 are RNA-1's aunts (children of NHP-1), not its ancestors.
    assert _uuids(resp) == ["D.SEQ-1", "NHP-1", "RNA-1", "TIS-1"]
    assert body["total_children"] == 3 and body["failed_uids"] == 0 and body["lineage_complete"] is True


def test_the_walk_passes_through_an_orphan_node(seek, graph):
    resp = _post({"identifiers": ["SLD-1"]}, login=SUPER)
    # SLD-1 -> ORPH-1 -> NHP-1 and down again; ORPH-1 itself has no MySQL row, so it is not in the download.
    assert "NHP-1" in _uuids(resp) and "ORPH-1" not in _uuids(resp)


def test_an_orphan_uid_that_mysql_still_holds_is_exported(seek, graph):
    with seek.cursor() as c:
        c.execute("INSERT INTO samples VALUES (500, 1, 'ORPH-1', '{}')")
    assert "ORPH-1" in _uuids(_post({"identifiers": ["SLD-1"]}, login=SUPER))


def test_include_tree_defaults_to_true(seek, graph):
    assert len(_uuids(_post({"identifiers": ["NHP-1"]}, login=SUPER))) > 1


def test_include_tree_false_never_walks_the_graph(seek, graph):
    resp = _post({"identifiers": ["NHP-1"], "include_tree": False}, login=SUPER)
    assert _uuids(resp) == ["NHP-1"] and graph.walked() == [] and resp.json()["lineage_complete"] is True


def test_a_sample_the_graph_has_not_reached_is_exported_not_404(seek, graph):
    """The read-after-write bug: an upload downloaded before the sync reached it answered 404."""
    resp = _post({"identifiers": ["TIS-NEW-1"]}, login=SUPER)
    assert resp.status_code == 200 and _uuids(resp) == ["TIS-NEW-1"]
    assert resp.json()["lineage_complete"] is False and resp["X-NExtSEEK-Lineage-Complete"] == "false"


def test_a_partial_graph_miss_keeps_every_requested_sample(seek, graph):
    resp = _post({"identifiers": ["TIS-NEW-1", "TIS-2"]}, login=SUPER)
    assert {"TIS-NEW-1", "TIS-2", "NHP-1"} <= set(_uuids(resp))
    assert resp.json()["failed_uids"] == 0 and resp.json()["lineage_complete"] is False


@pytest.mark.parametrize("login", [SUPER, MEMBER])
def test_a_graph_that_is_down_exports_the_requested_samples_and_says_so(seek, graph, login):
    graph.down = True
    resp = _post({"identifiers": ["TIS-2", "3"]}, login=login)
    assert resp.status_code == 200 and _uuids(resp) == ["TIS-2"]
    assert resp.json()["lineage_complete"] is False


def test_a_stale_graph_id_is_caught_and_the_uid_read_from_mysql(seek, graph):
    graph.nodes["TIS-2"] = (1, True)  # the graph maps TIS-2 to NHP-1's row
    resp = _post({"identifiers": ["TIS-2"], "include_tree": False}, login=SUPER)
    assert _uuids(resp) == ["TIS-2"]


def test_a_relative_whose_graph_id_is_stale_is_read_by_uuid(seek, graph):
    graph.nodes["TIS-2"] = (777, True)  # no row 777: the id moved after a delete and re-upload
    assert "TIS-2" in _uuids(_post({"identifiers": ["NHP-1"]}, login=SUPER))


def test_large_requests_are_read_in_chunks(seek, graph, monkeypatch):
    monkeypatch.setattr(sr, "MAX_IDS_PER_STATEMENT", 2)
    resp = _post({"identifiers": ["NHP-1"]}, login=SUPER)
    assert _uuids(resp) == ["D.SEQ-1", "DNA-OFF-1", "NHP-1", "RNA-1", "SLD-1", "TIS-1", "TIS-2", "TIS-FOR-1"]


# --------------------------------------------------------------------------- scope (#74)


def test_a_member_sees_only_their_projects(seek, graph):
    resp = _post({"identifiers": ["TIS-1"]})
    # DNA-OFF-1 is TIS-1's child but in project 4: dropped. RNA-1 is in two of the member's projects: once.
    assert _uuids(resp) == ["D.SEQ-1", "NHP-1", "RNA-1", "TIS-1"]
    assert resp.json()["total_samples"] == 4


def test_is_staff_does_not_widen_scope(seek, graph):
    staff = _client(MEMBER, staff=True).post(NEW, data={"identifiers": ["TIS-1"]}, format="json")
    plain = _client(MEMBER, staff=False).post(NEW, data={"identifiers": ["TIS-1"]}, format="json")
    assert staff.content == plain.content and "DNA-OFF-1" not in _uuids(staff)


@pytest.mark.parametrize("foreign, unknown", [("TIS-FOR-1", "TIS-NOPE"), ("6", "99")], ids=["uid", "seek-id"])
def test_a_foreign_identifier_answers_as_an_unknown_one(seek, graph, foreign, unknown):
    a, b = _post({"identifiers": [foreign]}), _post({"identifiers": [unknown]})
    assert (a.status_code, a.content) == (b.status_code, b.content) and a.status_code == 404
    assert "TIS-FOR-1" not in graph.walked()


def test_a_mixed_request_answers_as_if_the_foreign_uid_were_unknown(seek, graph):
    a, b = _post({"identifiers": ["TIS-2", "TIS-FOR-1"]}), _post({"identifiers": ["TIS-2", "TIS-NOPE"]})
    assert a.content == b.content and a.json()["failed_uids"] == 1
    assert "TIS-FOR-1" not in graph.walked()


def test_a_superuser_is_unscoped(seek, graph):
    resp = _post({"identifiers": ["TIS-FOR-1"]}, login=SUPER)
    assert set(_uuids(resp)) == {"TIS-FOR-1", "NHP-1"}


def test_a_member_of_no_project_sees_nothing_and_reads_no_graph(seek, graph):
    resp = _post({"identifiers": ["TIS-1"]}, login=NOBODY)
    assert resp.status_code == 404 and graph.walked() == []


def test_a_login_with_no_seek_person_sees_nothing(seek, graph):
    assert _post({"identifiers": ["TIS-1"]}, login="ghost").status_code == 404


def test_scope_that_cannot_be_read_fails_closed(seek, graph, monkeypatch):
    def broken(user):
        raise RuntimeError("db down")

    monkeypatch.setattr(sr, "resolve_scope", broken)
    assert _post({"identifiers": ["TIS-1"]}).status_code == 404


# --------------------------------------------------------------------------- the workbook


def test_the_workbook_download(seek, graph):
    resp = _post({"identifiers": ["TIS-2"], "output_format": "excel"}, login=SUPER)
    assert resp.status_code == 200
    assert resp["Content-Type"] == "application/vnd.ms-excel"
    assert resp["Content-Disposition"].startswith('attachment; filename="download-samples-')
    assert resp["X-NExtSEEK-Lineage-Complete"] == "true"
    book = load_workbook(BytesIO(b"".join(resp.streaming_content)))
    assert book.sheetnames[0] == "README" and {"NHP", "TIS"} <= set(book.sheetnames)
    assert book["README"]["A2"].value is None


def test_an_incomplete_lineage_is_flagged_in_the_workbook(seek, graph):
    graph.down = True
    resp = _post({"identifiers": ["TIS-2"], "output_format": "excel"}, login=SUPER)
    book = load_workbook(BytesIO(b"".join(resp.streaming_content)))
    assert resp["X-NExtSEEK-Lineage-Complete"] == "false"
    assert book["README"]["A2"].value == sr.LINEAGE_NOTICE


def test_the_browser_helper_posts_to_the_new_path():
    """Every download button goes through static/js/ns_sample_download.js (test_download_call_sites.py)."""
    import re
    from pathlib import Path

    from django.urls import resolve

    js = (Path(__file__).resolve().parents[2] / "static" / "js" / "ns_sample_download.js").read_text(encoding="utf-8")
    (endpoint,) = re.findall(r'var ENDPOINT = "([^"]+)";', js)
    assert endpoint == NEW and resolve(endpoint).func.cls.__name__ == "SampleRetrieveViewSet"


def test_a_uid_on_two_rows_exports_both_as_before(seek, graph):
    """A few UIDs sit on two MySQL rows. The old read by uuid returned both; the walk reaches one node."""
    with seek.cursor() as c:
        c.execute("INSERT INTO samples VALUES (600, 7, 'TIS-2', '{\"copy\": true}')")
    graph.twins["TIS-2"] = (600, True)
    ids = sorted(s["id"] for g in _post({"identifiers": ["NHP-1"]}, login=SUPER).json()["data"] for s in g["samples"])
    assert "600" in ids and "3" in ids


# --------------------------------------------------------------------------- review follow-ups


def test_a_members_uids_are_resolved_by_one_scoped_statement_never_the_graph_or_a_full_scan(seek, graph, monkeypatch):
    """A foreign UID and an unknown one must cost the same: resolving through the graph answered a foreign UID
    quickly and an unknown one through a full-table scan, and the time difference confirmed the foreign one exists."""
    scans = []
    real = sr._uuids_to_ids
    monkeypatch.setattr(sr, "_uuids_to_ids", lambda uuids: scans.append(sorted(uuids)) or real(uuids))
    for uid in ("TIS-FOR-1", "TIS-NOPE", "TIS-2"):
        _post({"identifiers": [uid], "include_tree": False})
    assert scans == [] and [c for c in graph.calls if c[0] == "resolve"] == []


def test_a_member_sees_both_rows_of_a_uid_on_two_rows(seek, graph):
    with seek.cursor() as c:
        c.execute("INSERT INTO samples VALUES (600, 7, 'TIS-2', '{}')")
        c.execute("INSERT INTO projects_samples VALUES (2, 600)")
    ids = sorted(s["id"] for g in _post({"identifiers": ["TIS-2"], "include_tree": False}).json()["data"]
                 for s in g["samples"])
    assert ids == ["3", "600"]


def test_a_non_ascii_digit_is_a_failed_identifier_not_a_500(seek, graph):
    resp = _post({"identifiers": ["TIS-2", "²"], "include_tree": False}, login=SUPER)
    assert resp.status_code == 200 and _uuids(resp) == ["TIS-2"] and resp.json()["failed_uids"] == 1


def test_a_rejected_statement_is_logged_as_an_error_not_a_blip(monkeypatch, caplog):
    from neo4j.exceptions import CypherSyntaxError

    class Driver:
        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def execute_query(self, *a, **k):
            raise CypherSyntaxError("Invalid input 'CYPHER 25'")

    monkeypatch.setattr(sr.GraphDatabase, "driver", lambda *a, **k: Driver())
    with caplog.at_level("WARNING", logger=sr.log.name), pytest.raises(sr.GraphUnavailable):
        sr._neo4j_run(sr.LINEAGE_CYPHER, uuids=["NHP-1"])
    assert [r.levelname for r in caplog.records] == ["ERROR"]


def test_the_graph_being_down_is_a_warning(monkeypatch, caplog):
    from neo4j.exceptions import ServiceUnavailable

    def down(*a, **k):
        raise ServiceUnavailable("connection refused")

    monkeypatch.setattr(sr.GraphDatabase, "driver", down)
    with caplog.at_level("WARNING", logger=sr.log.name), pytest.raises(sr.GraphUnavailable):
        sr._neo4j_run(sr.LINEAGE_CYPHER, uuids=["NHP-1"])
    assert [r.levelname for r in caplog.records] == ["WARNING"]
