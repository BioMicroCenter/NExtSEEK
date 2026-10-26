"""SEC-H03: user values never become formulas in exports."""
import io
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from nextseek_api.assistant.excel_export import generate_table_xlsx
from nextseek_api.services.sample_workbook import write_samples_workbook
from nextseek_api.services.sampletype_connections import rows_to_csv

EVIL = '=HYPERLINK("http://x","y")'
_MOD = "nextseek_api.services.sample_workbook"


def _assert_text(cell):
    assert cell.data_type == "s"
    assert cell.value == EVIL


def test_samples_workbook_pins_formula_as_text(tmp_path):
    df = pd.DataFrame([{"uuid": "MUS-230101ABC-1", "Name": EVIL, "Dose": "-80", "Note": "plain"}])
    out = tmp_path / "w.xlsx"
    ctx = {"MUS": {"name": "Mouse", "description": "d"}}
    with patch(f"{_MOD}.load_derivation_hops", return_value=[]), \
            patch(f"{_MOD}.load_sample_type_context", return_value=ctx):
        write_samples_workbook(df, str(out))
    ws = load_workbook(out)["MUS"]
    _assert_text(ws["A2"])
    assert ws["C2"].value == "plain" and ws["B2"].value == "-80"


def test_table_xlsx_pins_formula_as_text_and_keeps_numbers():
    data = generate_table_xlsx([{"label": "T", "columns": ["a", "n", "s"],
                                 "data": [{"a": EVIL, "n": -80, "s": "plain"}]}])
    ws = load_workbook(io.BytesIO(data))["T"]
    _assert_text(ws["A2"])
    assert ws["B2"].value == -80 and ws["B2"].data_type == "n"
    assert ws["C2"].value == "plain"


def test_nhp_workbook_pins_formula_as_text():
    from seek.timeline.services.nhp_service import save_nhp_data
    data = save_nhp_data([{"uuid": "MUS-1", "parsed_metadata": {"Name": EVIL, "Dose": -80}}])
    ws = load_workbook(io.BytesIO(data))["MUS"]
    _assert_text(ws["A2"])
    assert ws["B2"].value == -80


def test_connections_csv_prefixes_formula_only():
    rows = [{"parent_sample_type": EVIL, "child_sample_type": "-80", "internal_assay": "+1.5e3 ok",
             "n_edges": 3}]
    line = rows_to_csv(rows).splitlines()[1]
    assert line.startswith("\"'=HYPERLINK") and ",-80," in line
    assert "'+1.5e3 ok" in line
